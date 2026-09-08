#!/usr/bin/env python3
"""
Build per-pair apo/holo CS-list sequence alignment figures, a PPTX deck, and a
companion metadata CSV for a CSP_UBQ-style input CSV.

Alignment matches the pipeline: BMRB Seq_ID-offset pairing via
``scripts.csp._best_seqid_alignment`` / ``align_by_seqid_offset`` (exact AA
matches only). Figures are rendered with biotite's similarity-based alignment
plotter from those gapped strings.

Outputs (default under outputs/apo_holo_cs_list_alignments/):
  figures/{HOLO_PDB}_{apo_bmrb}.png
  apo_holo_cs_list_alignments.pptx
  apo_holo_cs_list_metadata.csv
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import matplotlib
import numpy as np

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import biotite.sequence as seq  # noqa: E402
import biotite.sequence.align as align  # noqa: E402
import biotite.sequence.graphics as graphics  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

_REPO = Path(__file__).resolve().parent.parent
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from scripts.apo_holo_exp_conditions import (  # noqa: E402
    ExtractedConditions,
    ensure_star_path,
    extract_conditions_from_star,
)
from scripts.apo_polymer_inventory import inventory_entities  # noqa: E402
from scripts.bmrb_io import parse_sequence_and_shifts_from_saveframes  # noqa: E402
from scripts.config import Paths  # noqa: E402
from scripts.csp import _best_seqid_alignment  # noqa: E402
from scripts.target_resolution import canonical_output_dir_name  # noqa: E402

_PROTEIN_ALPHABET = set(seq.ProteinSequence.alphabet)
_SYMBOLS_PER_LINE = 60

CSV_FIELDS = [
    "apo_bmrb",
    "holo_bmrb",
    "apo_pdb",
    "holo_pdb",
    "apo_saveframe",
    "holo_saveframe",
    "seqid_offset",
    "apo_pH",
    "apo_temperature_C",
    "apo_pressure",
    "apo_pressure_units",
    "holo_pH",
    "holo_temperature_C",
    "holo_pressure",
    "holo_pressure_units",
    "apo_entities",
    "holo_entities",
    "alignment_score",
    "identity_pct",
    "n_mapped_pairs",
    "figure_png",
    "Notes",
]


def _repo_root() -> Path:
    return _REPO


def _fmt_num(val: Optional[float], digits: int) -> str:
    if val is None:
        return ""
    return f"{val:.{digits}f}".rstrip("0").rstrip(".")


def _sanitize_protein_seq(raw: str) -> str:
    """Map non-alphabet residue codes to X so biotite ProteinSequence accepts them."""
    out: List[str] = []
    for ch in (raw or "").upper():
        if ch in ("-", ".", " ", "\t"):
            continue
        out.append(ch if ch in _PROTEIN_ALPHABET else "X")
    return "".join(out)


def entity_names_joined(star_path: Path) -> str:
    """Join entity display names in STAR inventory order (polymers then ligands)."""
    _fmt, entities = inventory_entities(star_path)
    # Prefer numeric Entity.ID order when present (NMR-STAR 3); else keep inventory order.
    if any((e.entity_id or "").strip() for e in entities):
        def _sort_key(e) -> Tuple[int, int]:
            try:
                eid = int(float(e.entity_id)) if e.entity_id else 10**9
            except (TypeError, ValueError):
                eid = 10**9
            return (eid, 0)

        ordered = sorted(enumerate(entities), key=lambda iv: (_sort_key(iv[1])[0], iv[0]))
        entities = [e for _, e in ordered]

    names: List[str] = []
    seen = set()
    for e in entities:
        name = (e.name or "").strip()
        if not name or name in (".", "?"):
            continue
        if name in seen:
            continue
        seen.add(name)
        names.append(name)
    return "; ".join(names)

def _identity_and_mapped(gapped_a: str, gapped_b: str) -> Tuple[float, int]:
    matches = 0
    compared = 0
    mapped = 0
    for a, b in zip(gapped_a, gapped_b):
        if a != "-" and b != "-":
            mapped += 1
            compared += 1
            if a == b:
                matches += 1
    identity = (matches / compared * 100.0) if compared else 0.0
    return identity, mapped


def biotite_alignment_from_gapped(
    gapped_a: str, gapped_b: str, score: float
) -> Tuple[align.Alignment, align.SubstitutionMatrix]:
    """Rebuild a biotite Alignment from Seq_ID-offset gapped strings for plotting."""
    if len(gapped_a) != len(gapped_b):
        raise ValueError("gapped sequences must have equal length")
    # Sanitize per column first so ungapped lengths match the trace walk
    # (preserves '*'; maps unknown symbols to 'X').
    g_a = "".join("-" if c == "-" else (_sanitize_protein_seq(c) or "X") for c in gapped_a)
    g_b = "".join("-" if c == "-" else (_sanitize_protein_seq(c) or "X") for c in gapped_b)
    ungapped_a = "".join(c for c in g_a if c != "-")
    ungapped_b = "".join(c for c in g_b if c != "-")
    if not ungapped_a or not ungapped_b:
        raise RuntimeError("empty sequence after removing gaps")
    sa = seq.ProteinSequence(ungapped_a)
    sb = seq.ProteinSequence(ungapped_b)
    trace: List[List[int]] = []
    ia = ib = 0
    for ca, cb in zip(g_a, g_b):
        if ca == "-" and cb == "-":
            # Both-gap columns can appear after X/X conflicts; skip for biotite.
            continue
        if ca != "-" and cb != "-":
            trace.append([ia, ib])
            ia += 1
            ib += 1
        elif ca == "-":
            trace.append([-1, ib])
            ib += 1
        else:
            trace.append([ia, -1])
            ia += 1
    if ia != len(sa) or ib != len(sb):
        raise RuntimeError(
            f"trace length mismatch: ia={ia}/{len(sa)} ib={ib}/{len(sb)}"
        )
    matrix = align.SubstitutionMatrix.std_protein_matrix()
    aln = align.Alignment([sa, sb], np.asarray(trace, dtype=np.int64), float(score))
    return aln, matrix


@dataclass
class PairAlignment:
    alignment: align.Alignment
    matrix: align.SubstitutionMatrix
    apo_saveframe: str
    holo_saveframe: str
    score: float
    identity_pct: float
    n_mapped_pairs: int
    apo_seq: str
    holo_seq: str
    seqid_offset: int
    aligned_apo: str
    aligned_holo: str


def best_seqid_pair_alignment(
    apo_sequences: Sequence,
    holo_sequences: Sequence,
) -> PairAlignment:
    """Select polymers and align exactly as the CSP pipeline does."""
    best = _best_seqid_alignment(apo_sequences, holo_sequences)
    if best is None or not best["mapping"]:
        raise RuntimeError("no Seq_ID-offset exact AA matches between apo/holo CS lists")
    g_a = best["aligned_apo"]
    g_b = best["aligned_holo"]
    identity, n_mapped = _identity_and_mapped(g_a, g_b)
    aln, matrix = biotite_alignment_from_gapped(g_a, g_b, float(best["score"]))
    return PairAlignment(
        alignment=aln,
        matrix=matrix,
        apo_saveframe=best["apo_saveframe"],
        holo_saveframe=best["holo_saveframe"],
        score=float(best["score"]),
        identity_pct=identity,
        n_mapped_pairs=n_mapped,
        apo_seq=best["apo_seq"],
        holo_seq=best["holo_seq"],
        seqid_offset=int(best["offset"]),
        aligned_apo=g_a,
        aligned_holo=g_b,
    )

def _conditions_lines(label: str, bmrb: str, sf: str, ec: ExtractedConditions, entities: str) -> List[str]:
    press = ""
    if ec.pressure is not None:
        units = ec.pressure_units or ""
        press = f"  pressure={_fmt_num(ec.pressure, 3)}{(' ' + units) if units else ''}"
    ph = _fmt_num(ec.pH, 3) or "n/a"
    temp = _fmt_num(ec.temperature_C, 2) or "n/a"
    ent = entities if entities else "(none)"
    return [
        f"{label} BMRB {bmrb}  saveframe={sf}",
        f"  pH={ph}  T={temp} °C{press}",
        f"  Entities: {ent}",
    ]


def render_alignment_png(
    out_png: Path,
    *,
    holo_pdb: str,
    apo_bmrb: str,
    holo_bmrb: str,
    pair: PairAlignment,
    apo_ec: ExtractedConditions,
    holo_ec: ExtractedConditions,
    apo_entities: str,
    holo_entities: str,
) -> None:
    out_png.parent.mkdir(parents=True, exist_ok=True)

    header_lines = [
        f"{holo_pdb.upper()}  apo={apo_bmrb}  holo={holo_bmrb}",
        f"Method=Seq_ID offset  offset={pair.seqid_offset}  "
        f"score={pair.score:.1f}  identity={pair.identity_pct:.1f}%  "
        f"mapped={pair.n_mapped_pairs}",
        "",
        *_conditions_lines("Apo", apo_bmrb, pair.apo_saveframe, apo_ec, apo_entities),
        "",
        *_conditions_lines("Holo", holo_bmrb, pair.holo_saveframe, holo_ec, holo_entities),
    ]
    # Wrap long entity lines for the figure header
    wrapped: List[str] = []
    for ln in header_lines:
        if ln.startswith("  Entities:") and len(ln) > 110:
            wrapped.extend(textwrap.wrap(ln, width=110, subsequent_indent="    "))
        else:
            wrapped.append(ln)

    aln_len = len(pair.alignment)
    n_blocks = max(1, math.ceil(aln_len / _SYMBOLS_PER_LINE))
    header_h = 0.28 * len(wrapped) + 0.35
    align_h = max(2.2, 1.15 * n_blocks + 0.6)
    fig_h = header_h + align_h
    fig_w = 12.5

    fig = plt.figure(figsize=(fig_w, fig_h), dpi=150)
    gs = fig.add_gridspec(2, 1, height_ratios=[header_h, align_h], hspace=0.08)
    ax_hdr = fig.add_subplot(gs[0])
    ax_aln = fig.add_subplot(gs[1])

    ax_hdr.axis("off")
    ax_hdr.set_xlim(0, 1)
    ax_hdr.set_ylim(0, 1)
    y0 = 1.0
    dy = 1.0 / max(len(wrapped), 1)
    for i, ln in enumerate(wrapped):
        weight = "bold" if i == 0 else "normal"
        ax_hdr.text(
            0.0,
            y0 - (i + 0.7) * dy,
            ln,
            fontsize=8 if i else 10,
            fontfamily="monospace",
            fontweight=weight,
            va="top",
            ha="left",
            transform=ax_hdr.transAxes,
        )

    graphics.plot_alignment_similarity_based(
        ax_aln,
        pair.alignment,
        matrix=pair.matrix,
        labels=["Apo", "Holo"],
        show_numbers=True,
        show_line_position=True,
        symbols_per_line=_SYMBOLS_PER_LINE,
    )

    fig.savefig(out_png, dpi=150, bbox_inches="tight", pad_inches=0.25)
    plt.close(fig)


def add_image_slide(prs: Presentation, png_path: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    im = Image.open(png_path).convert("RGB")
    iw, ih = im.size
    sw = float(prs.slide_width)
    sh = float(prs.slide_height)
    aspect_img = iw / ih
    aspect_slide = sw / sh
    if aspect_img > aspect_slide:
        width = sw
        height = sw / aspect_img
    else:
        height = sh
        width = sh * aspect_img
    left = int((sw - width) / 2)
    top = int((sh - height) / 2)
    slide.shapes.add_picture(str(png_path), left, top, width=int(width), height=int(height))


def add_error_slide(prs: Presentation, title: str, message: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.3), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.bold = True
    p2 = tf.add_paragraph()
    p2.text = message
    p2.font.size = Pt(16)


def _empty_meta_row(row: Dict[str, str], error: str = "") -> Dict[str, str]:
    return {
        "apo_bmrb": (row.get("apo_bmrb") or "").strip(),
        "holo_bmrb": (row.get("holo_bmrb") or "").strip(),
        "apo_pdb": (row.get("apo_pdb") or "").strip(),
        "holo_pdb": (row.get("holo_pdb") or "").strip(),
        "apo_saveframe": "",
        "holo_saveframe": "",
        "seqid_offset": "",
        "apo_pH": "",
        "apo_temperature_C": "",
        "apo_pressure": "",
        "apo_pressure_units": "",
        "holo_pH": "",
        "holo_temperature_C": "",
        "holo_pressure": "",
        "holo_pressure_units": "",
        "apo_entities": "",
        "holo_entities": "",
        "alignment_score": "",
        "identity_pct": "",
        "n_mapped_pairs": "",
        "figure_png": "",
        "Notes": error,
    }


def process_pair(
    row: Dict[str, str],
    *,
    cs_dir: Path,
    figures_dir: Path,
    do_fetch: bool,
) -> Tuple[Dict[str, str], Optional[Path], Optional[str]]:
    apo_bmrb = (row.get("apo_bmrb") or "").strip()
    holo_bmrb = (row.get("holo_bmrb") or "").strip()
    apo_pdb = (row.get("apo_pdb") or "").strip()
    holo_pdb = (row.get("holo_pdb") or "").strip()
    if not apo_bmrb or not holo_bmrb or not holo_pdb:
        return _empty_meta_row(row, "missing apo_bmrb/holo_bmrb/holo_pdb"), None, "missing IDs"

    apo_star = ensure_star_path(apo_bmrb, cs_dir, do_fetch)
    holo_star = ensure_star_path(holo_bmrb, cs_dir, do_fetch)
    if apo_star is None or holo_star is None:
        missing = []
        if apo_star is None:
            missing.append(f"apo STAR {apo_bmrb}")
        if holo_star is None:
            missing.append(f"holo STAR {holo_bmrb}")
        err = "missing " + ", ".join(missing)
        return _empty_meta_row(row, err), None, err

    apo_sequences = parse_sequence_and_shifts_from_saveframes(str(apo_star))
    holo_sequences = parse_sequence_and_shifts_from_saveframes(str(holo_star))
    if not apo_sequences or not holo_sequences:
        err = "no H+N CS-list saveframes"
        return _empty_meta_row(row, err), None, err

    pair = best_seqid_pair_alignment(apo_sequences, holo_sequences)
    apo_ec = extract_conditions_from_star(str(apo_star))
    holo_ec = extract_conditions_from_star(str(holo_star))
    apo_entities = entity_names_joined(apo_star)
    holo_entities = entity_names_joined(holo_star)

    stem = canonical_output_dir_name(holo_pdb, apo_bmrb)
    out_png = figures_dir / f"{stem}.png"
    render_alignment_png(
        out_png,
        holo_pdb=holo_pdb,
        apo_bmrb=apo_bmrb,
        holo_bmrb=holo_bmrb,
        pair=pair,
        apo_ec=apo_ec,
        holo_ec=holo_ec,
        apo_entities=apo_entities,
        holo_entities=holo_entities,
    )

    meta = {
        "apo_bmrb": apo_bmrb,
        "holo_bmrb": holo_bmrb,
        "apo_pdb": apo_pdb,
        "holo_pdb": holo_pdb,
        "apo_saveframe": pair.apo_saveframe,
        "holo_saveframe": pair.holo_saveframe,
        "seqid_offset": str(pair.seqid_offset),
        "apo_pH": _fmt_num(apo_ec.pH, 3),
        "apo_temperature_C": _fmt_num(apo_ec.temperature_C, 2),
        "apo_pressure": _fmt_num(apo_ec.pressure, 3),
        "apo_pressure_units": apo_ec.pressure_units or "",
        "holo_pH": _fmt_num(holo_ec.pH, 3),
        "holo_temperature_C": _fmt_num(holo_ec.temperature_C, 2),
        "holo_pressure": _fmt_num(holo_ec.pressure, 3),
        "holo_pressure_units": holo_ec.pressure_units or "",
        "apo_entities": apo_entities,
        "holo_entities": holo_entities,
        "alignment_score": f"{pair.score:.2f}",
        "identity_pct": f"{pair.identity_pct:.1f}",
        "n_mapped_pairs": str(pair.n_mapped_pairs),
        "figure_png": str(out_png.relative_to(_repo_root())) if out_png.is_relative_to(_repo_root()) else str(out_png),
        "Notes": "",
    }
    return meta, out_png, None

def main() -> None:
    paths = Paths()
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--csv",
        type=Path,
        default=_repo_root() / "data" / "CSP_UBQ_ph0.5_temp5C.csv",
        help="Input pair CSV (apo_bmrb, holo_bmrb, holo_pdb, ...)",
    )
    ap.add_argument(
        "--out-dir",
        type=Path,
        default=_repo_root() / "outputs" / "apo_holo_cs_list_alignments",
        help="Output directory for figures/, PPTX, and CSV",
    )
    ap.add_argument(
        "--cs-dir",
        type=Path,
        default=Path(paths.cs_cache_dir),
        help="Directory of cached BMRB STAR files",
    )
    ap.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Process only the first N rows (0 = all)",
    )
    ap.add_argument(
        "--no-fetch",
        action="store_true",
        help="Do not download missing STAR files",
    )
    args = ap.parse_args()

    out_dir: Path = args.out_dir
    figures_dir = out_dir / "figures"
    figures_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = out_dir / "apo_holo_cs_list_alignments.pptx"
    csv_path = out_dir / "apo_holo_cs_list_metadata.csv"

    with args.csv.open(newline="", encoding="utf-8") as f:
        rows: List[Dict[str, str]] = list(csv.DictReader(f))
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    meta_rows: List[Dict[str, str]] = []
    n_ok = 0
    n_err = 0
    do_fetch = not args.no_fetch

    for i, row in enumerate(rows, start=1):
        apo = (row.get("apo_bmrb") or "").strip()
        holo_pdb = (row.get("holo_pdb") or "").strip()
        label = f"{holo_pdb}_{apo}" if holo_pdb and apo else f"row_{i}"
        try:
            meta, png, err = process_pair(
                row,
                cs_dir=args.cs_dir,
                figures_dir=figures_dir,
                do_fetch=do_fetch,
            )
        except Exception as exc:  # noqa: BLE001 — keep deck complete
            err = str(exc)
            meta = _empty_meta_row(row, err)
            png = None
        meta_rows.append(meta)
        if png is not None and png.is_file():
            add_image_slide(prs, png)
            n_ok += 1
            print(f"[{i}/{len(rows)}] OK  {label} -> {png.name}")
        else:
            add_error_slide(prs, label, err or meta.get("Notes") or "failed")
            n_err += 1
            print(f"[{i}/{len(rows)}] ERR {label}: {err or meta.get('Notes')}")

    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(meta_rows)

    prs.save(pptx_path)
    print(
        f"Wrote {len(meta_rows)} CSV rows, {len(prs.slides)} slides "
        f"({n_ok} ok, {n_err} err) -> {out_dir.resolve()}"
    )


if __name__ == "__main__":
    main()
